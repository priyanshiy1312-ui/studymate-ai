import re, uuid, os, shutil
from pathlib import Path
from llm import chat_json, describe_image

UP = Path(__file__).parent / "data" / "uploads"
UP.mkdir(parents=True, exist_ok=True)
TS = re.compile(r"^\[?(?:(\d+):)?(\d{1,2}):(\d{2})\]?\s*(.*)")

def mk(sid, name, typ, loc, text, topic="General"):
    return dict(id=uuid.uuid4().hex[:12], source_id=sid, source_name=name, type=typ, loc=loc, text=text.strip(), topic=topic)

def words_chunks(text, size=160, overlap=25):
    w, out, i = text.split(), [], 0
    while i < len(w):
        out.append(" ".join(w[i:i + size])); i += size - overlap
    return [c for c in out if len(c.split()) >= 6]

def ingest_pdf(path, sid, name, furl):
    from pypdf import PdfReader
    out = []
    for n, p in enumerate(PdfReader(str(path)).pages, 1):
        loc = {"page": n, "file": furl}
        for c in words_chunks(p.extract_text() or ""):
            out.append(mk(sid, name, "pdf", loc, c))
        try:
            for im in list(p.images)[:2]:
                if len(im.data) > 5000:
                    mt = "image/jpeg" if im.name.lower().endswith(("jpg", "jpeg")) else "image/png"
                    out.append(mk(sid, name, "pdf", loc, "[Figure] " + describe_image(im.data, mt)))
        except Exception:
            pass
    return out

def ingest_pptx(path, sid, name, furl):
    from pptx import Presentation
    out = []
    for n, s in enumerate(Presentation(str(path)).slides, 1):
        loc = {"slide": n, "file": furl}
        title = s.shapes.title.text if s.shapes.title is not None and s.shapes.title.text.strip() else "General"
        text = "\n".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip())
        if s.has_notes_slide:
            text += "\n" + s.notes_slide.notes_text_frame.text
        if text.strip():
            out.append(mk(sid, name, "slide", loc, text, title))
        for sh in s.shapes:
            if sh.shape_type == 13:
                try:
                    out.append(mk(sid, name, "slide", loc, "[Diagram] " + describe_image(sh.image.blob, sh.image.content_type), title))
                except Exception:
                    pass
    return out

def group_segments(segs, sid, name, extra):
    out, buf, start = [], [], 0
    for sec, t in segs:
        if not buf: start = sec
        buf.append(t)
        if sec - start >= 60 and len(" ".join(buf).split()) >= 40:
            out.append(mk(sid, name, "video", {"t": int(start), **extra}, " ".join(buf))); buf = []
    if buf:
        out.append(mk(sid, name, "video", {"t": int(start), **extra}, " ".join(buf)))
    return out

def ingest_transcript(path, sid, name, furl):
    segs = []
    for line in open(path, encoding="utf8"):
        m = TS.match(line.strip())
        if m:
            h, mi, s, t = m.groups(); segs.append((int(h or 0) * 3600 + int(mi) * 60 + int(s), t))
    return group_segments(segs, sid, name, {})

def ingest_media(path, sid, name, extra):
    from faster_whisper import WhisperModel
    model = WhisperModel(os.getenv("WHISPER_MODEL", "base"), compute_type="int8")
    segs, _ = model.transcribe(str(path))
    return group_segments([(s.start, s.text.strip()) for s in segs], sid, name, extra)

def ingest_youtube(url):
    import yt_dlp
    sid = uuid.uuid4().hex[:8]
    with yt_dlp.YoutubeDL({"format": "bestaudio/best", "outtmpl": str(UP / f"{sid}.%(ext)s"), "quiet": True}) as y:
        info = y.extract_info(url, download=True); f = y.prepare_filename(info)
    chunks = ingest_media(f, sid, info.get("title", "YouTube lecture"), {"url": url})
    tag_topics(chunks)
    return chunks

def tag_topics(chunks):
    for i in range(0, len(chunks), 20):
        b = chunks[i:i + 20]
        try:
            r = chat_json("Give each numbered text a short topic name (1-4 words, Title Case). Reuse identical names for the same concept. Return a JSON list of strings in the same order.",
                          "\n".join(f"{j}. {c['text'][:250]}" for j, c in enumerate(b)))
            for c, t in zip(b, r):
                if isinstance(t, str) and t.strip(): c["topic"] = t.strip()
        except Exception:
            pass

def ingest_path(path, filename):
    sid = uuid.uuid4().hex[:8]
    dest = UP / f"{sid}_{filename}"
    shutil.copy(path, dest)
    furl, name, ext = f"/files/{dest.name}", Path(filename).stem, Path(filename).suffix.lower()
    if ext == ".pdf": chunks = ingest_pdf(dest, sid, name, furl)
    elif ext == ".pptx": chunks = ingest_pptx(dest, sid, name, furl)
    elif ext in (".txt", ".vtt"): chunks = ingest_transcript(dest, sid, name, furl)
    elif ext in (".mp4", ".mp3", ".wav", ".m4a", ".mkv", ".webm"): chunks = ingest_media(dest, sid, name, {"file": furl})
    else: raise ValueError(f"Unsupported file type: {ext}")
    tag_topics(chunks)
    return chunks
